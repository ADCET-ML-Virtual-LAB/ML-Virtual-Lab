"""Excel (.xlsx) export for the instructor console.

Layout: the details that are the same for every row (batch, academic year,
experiment number, experiment name) sit on top in merged cells; below them is a
clean table with a bold header row and sized columns.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_THIN = Side(style="thin", color="BBBBBB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="DBEAFE")
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _set(cell, value):
    cell.value = value
    if isinstance(value, str):
        # Always store text as text, so a name starting with "=" is never treated as a formula.
        cell.data_type = "s"


def build_sheet(
    title: str,
    details: list[tuple[str, str]],
    headers: list[str],
    rows: list[list],
    widths: list[float],
) -> bytes:
    """Return an .xlsx file as bytes.

    details: (label, value) pairs, each written as "label: value" in one merged row
    headers/rows/widths: the table; widths are column widths in characters.
    A numeric cell is shown with a trailing % sign (the value stays a number).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = title
    last_col = len(headers)

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = 1
    for label, value in details:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        cell = ws.cell(row=row, column=1)
        _set(cell, f"{label}: {value}")
        cell.font = Font(bold=True, size=12)
        cell.alignment = _LEFT
        ws.row_dimensions[row].height = 24
        row += 1

    row += 1  # blank row between the details and the table
    header_row = row
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col)
        _set(cell, text)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = _LEFT
    ws.row_dimensions[row].height = 22

    for values in rows:
        row += 1
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col)
            _set(cell, value)
            cell.border = _BORDER
            cell.alignment = _LEFT
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = 'General"%"'

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
